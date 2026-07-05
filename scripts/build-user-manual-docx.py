from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "docs/McLellan-Hub-User-Manual.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                set_cell_width(row.cells[idx], width)
                row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(item)


def add_table(doc, headers, rows, widths, header_fill="E8EEF5", font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, h in enumerate(headers):
        set_cell_shading(hdr[idx], header_fill)
        set_cell_margins(hdr[idx])
        r = hdr[idx].paragraphs[0].add_run(h)
        set_run_font(r, bold=True, size=font_size)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, bold=True, color="1F4D78")
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    p2.add_run(body)
    set_table_width(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "McLellan Hub User Manual"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("666666")
    footer = section.footer.paragraphs[0]
    footer.text = "Private user reference - current as of 5 July 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor.from_string("666666")


def page_inventory():
    return [
        ("/", "Hub Home", "Start page with recent chats, active projects, calendar context, token/newsletter signals, and shortcuts.", "Use it to orient yourself at the start of a session."),
        ("/c", "AI Chat", "General assistant workspace with model selection, conversation history, uploads, export, ratings, recall, and project context.", "Use for research, drafting, analysis, task commands, and reusable project work."),
        ("/p/:slug", "Project Chat", "A chat scoped to one project, using its saved messages and uploaded documents as context.", "Use when the answer should stay attached to a continuing workstream."),
        ("/settings", "Hub Settings", "User-level model defaults and available chat models.", "Use when changing your normal chat model preference."),
        ("/logs", "Chat Logs", "Recent conversations and assistant outputs.", "Use to retrieve or inspect past work."),
        ("/humanizer", "AI Humanizer", "Redrafts AI prose and reports what machine-writing tells were changed.", "Use before publishing or sending text that must sound less generic."),
        ("/crm", "CRM Overview", "Relationship dashboard across people, companies, meetings, tasks, reminders, projects, reports, and Ask the Hub.", "Use as the relationship command center."),
        ("/crm/contacts", "People", "List and create people records.", "Use to find or add a contact."),
        ("/crm/contact/:id", "Person Page", "Timeline, tasks, companies, projects, knowledge claims, meetings, edit/merge/delete controls.", "Use before or after interacting with someone."),
        ("/crm/companies", "Companies", "List and create organisations.", "Use to manage organisations and linked people."),
        ("/crm/company/:id", "Company Page", "Company details, linked people, tasks, knowledge, meetings, and timeline.", "Use before supplier/client/employer conversations."),
        ("/crm/meetings", "Meetings", "Meeting list and manual meeting creation.", "Use to prepare for or review meetings."),
        ("/crm/meeting-intake", "Meeting Intake", "Upload or paste a transcript and let the Hub extract summary, attendees, actions, CRM updates, project notes, and open questions.", "Use after calls, Krisp exports, or meeting notes."),
        ("/crm/meeting/:id", "Meeting Detail", "Edit metadata, attendees, notes, and meeting-sourced facts.", "Use to correct or enrich a meeting record."),
        ("/crm/tasks", "Tasks", "Google Tasks backed task list with Hub links to people, companies, projects, sources, due date/time, history, and restore controls.", "Use as the daily task queue."),
        ("/crm/tasks/:id", "Task Detail", "Edit task title, notes, due date/time, CRM links, project links, and subtasks.", "Use when a task needs context or correction."),
        ("/crm/reminders", "Reminders", "Reminder queue with open, due, general, and resolved items.", "Use for time-based nudges that are not full tasks."),
        ("/crm/projects", "CRM Projects", "Project list with activity and linked context.", "Use to browse relationship-backed workstreams."),
        ("/crm/project/:slug", "Project Page", "Project briefing with contacts, companies, meetings, documents, tasks, facts, knowledge, and recent chat memory.", "Use to understand a workstream without hunting across tools."),
        ("/crm/project-report", "Project Report", "AI-written report generated from meetings, tasks, emails, documents, and knowledge atoms.", "Use for progress summaries and stakeholder updates."),
        ("/crm/knowledge", "Ask the Hub", "Free-form questions over knowledge atoms, insight atoms, and semantic search.", "Use when you want the Hub to answer from what it knows."),
        ("/crm/source/:kind/:id", "Source Viewer", "Shows the source behind a fact or knowledge claim.", "Use to check evidence."),
        ("/wiki", "Wiki Search", "Search the private knowledge base and get a source-backed answer.", "Use for durable knowledge retrieval."),
        ("/wiki/browse", "Wiki Browse", "Maintenance browse of indexed material by type.", "Use when search is not enough."),
        ("/wiki/page/new", "New Wiki Page", "Manual page creation.", "Use for durable notes that should live in the wiki."),
        ("/wiki/page/:slug", "Wiki Page", "Readable page with sources and related material.", "Use as the human-readable knowledge layer."),
        ("/wiki/graph", "Wiki Graph", "Visual graph of wiki links and manual links.", "Use to inspect relationships."),
        ("/wiki/orphans", "Wiki Orphans", "Pages with weak or missing links.", "Use to keep the knowledge base connected."),
        ("/wiki/ingest", "Add to Knowledge Base", "Queue YouTube or web URLs for Synthadoc/vault ingest.", "Use to add external material."),
        ("/newsletter", "Intelligence", "Newsletter topic review, briefing generation, send/publish/wiki/export actions, interests, formats, and schedules.", "Use to turn incoming newsletters into weekly briefings."),
        ("/newsletter/sources", "Newsletter Sources", "Source management and priority.", "Use to tune incoming intelligence."),
        ("/newsletter/creators", "Creator Feeds", "RSS/watchlist feed management and reading flows.", "Use to maintain followed creators and sites."),
        ("/lin", "Content Create", "LinkedIn topic research, draft, score, refine, carousel, image-prompt flow.", "Use to create a post package."),
        ("/lin/plan", "Content Plan", "Cadence policy, topic plans, suggestions, and day-by-day research.", "Use to plan content ahead."),
        ("/lin/queue", "Content Queue", "Processing, draft, scheduled, and error posts.", "Use to finish or fix work in progress."),
        ("/lin/published", "Published Content", "Published post archive with sorting and filters.", "Use to audit what has gone live."),
        ("/debrief", "Daily Debrief", "Voice-led end-of-day operational review.", "Use to capture what happened and what follows."),
        ("/flights", "Flights", "Flight history, manual import, file import, lookup, bulk lookup, and statistics.", "Use for travel records and flight task preparation."),
        ("/token-burn", "Token Burn", "Model/token usage reporting.", "Use to monitor cost and usage."),
        ("/prompt", "Prompt Library", "Prompt assets, builder, gym, inbox, and audit agent packs.", "Use to save and improve reusable prompts."),
        ("/admin", "Hub Admin", "Project memory and document administration.", "Use to manage projects, memory, and admin-only controls."),
    ]


def admin_inventory():
    return [
        ("/admin", "Project admin", "Create/edit/delete projects, set context depth, mark CV context, inspect project memory counts."),
        ("/admin/projects/:slug", "Project memory", "View/delete project memories and documents, import from Drive, manage wiki tags/pins."),
        ("/admin/crm", "CRM admin", "Admin view for contacts and CRM maintenance."),
        ("/admin/email-taxonomy", "Email taxonomy", "Maintain canonical Gmail/AgentMail classification labels and trigger email processing."),
        ("/admin/chatlogs", "Chat logs", "Inspect chat outputs by rating, age, and sort order."),
        ("/admin/newsletter-ingestion", "Newsletter ingestion", "Audit newsletter extraction status, errors, source data, and direct/fallback extraction."),
        ("/admin/debrief", "Debrief sessions", "Review debrief transcripts and extracted output."),
        ("/admin/models", "Models", "Add, edit, enable, disable, delete, test, and set default chat models."),
        ("/admin/models/system", "System models", "Assign models to background AI slots such as email classifier, meeting intake, knowledge query, and debrief voice."),
        ("/admin/models/prompts", "System prompts", "Edit or restore system prompt text for visible slots; shape prompts for a model family."),
        ("/admin/models/tiers", "Model tiers", "Group models in the chat picker and set default search mode per tier."),
        ("/admin/knowledge", "Knowledge review", "Review contradictions, stale claims, duplicate/supersession decisions, knowledge receipts, and atom status."),
        ("/admin/shortcuts", "Shortcuts", "Manage quick-start cards and model shortcuts."),
        ("/admin/test", "Model test arena", "Run prompts across selected models, upload a file, improve prompts, compare outputs, and save logs."),
        ("/admin/nakai-briefings", "Nakai briefings", "Build, resend, inspect, and manage daily regulatory briefings and reference sources."),
        ("/admin/linkedin", "LinkedIn admin", "Inspect posts, topics, statuses, schedules, and published title handling."),
        ("/admin/jobs", "Jobs", "Inspect pending and recent background jobs; cancel pending jobs."),
        ("/admin/connectivity", "Connectivity", "Find orphaned data and run the mycelium connection pass."),
    ]


def model_slots():
    return [
        ("Chat infrastructure", "recall_tagger", "Tags each chat turn for later recall.", "Shared", "Cheap model is enough."),
        ("Chat infrastructure", "multisearch_planner", "Plans multiple web searches for research mode.", "Shared", "Needs good query variety."),
        ("Chat infrastructure", "multisearch_synthesiser", "Writes the final answer from gathered sources.", "Shared", "Use stronger reasoning for source reconciliation."),
        ("Background processing", "crm_parser", "Interprets manually saved CRM notes.", "Shared", "Affects note-to-fact/reminder behavior."),
        ("Background processing", "email_classifier", "Classifies Gmail and stores summary/source evidence.", "Shared", "Affects labels and source summaries."),
        ("Background processing", "meeting_intake", "Extracts meeting summary, attendees, CRM updates, actions, project notes, open questions.", "Per user", "Use a strong long-context model."),
        ("Background processing", "agentmail_extractor", "Extracts people, facts, and actions from AgentMail.", "Shared", "Important for AI-facing inbox."),
        ("Background processing", "task_extractor", "Extracts follow-up tasks from documents and learns from rejected suggestions.", "Shared", "Feeds Google Tasks via review/knowledge path."),
        ("Background processing", "reg_synopsis", "Assesses regulatory publications for Nakai alerts.", "Shared", "Needs evidence discipline."),
        ("Background processing", "hub_dev_constraint", "Knowledge-first constraint used in Hub coding/prompt work.", "Shared", "Guardrail prompt."),
        ("Background processing", "prompt_improver", "Rewrites prompts in the admin test panel.", "Shared", "Fast editor slot."),
        ("Background processing", "prompt_shaper", "Restyles system prompts for a target model family.", "Shared", "Needs strong model; proposal only."),
        ("Background processing", "style_distiller", "Monthly distillation of model-family prompt style profiles.", "Shared", "Feeds prompt shaping."),
        ("Background processing", "prompt_adapter", "Builds reusable prompts from rough prompts and examples.", "Shared", "Prompt Library builder."),
        ("Background processing", "prompt_optimizer", "Optimises reusable prompts against examples and a rubric.", "Shared", "Prompt Gym."),
        ("Background processing", "opportunity_extractor", "Extracts short-lived offers/opportunities from inbound email.", "Shared", "Used by suggestion engine."),
        ("Background processing", "task_rule_learner", "Turns a wrongly created task into a reusable decision rule.", "Shared", "Prompt-only child of task extractor."),
        ("Background processing", "admin_synthesiser", "Synthesises admin test multi-search results.", "Shared", "Test arena support."),
        ("Suggestion engine", "suggestions", "Parent model for travel, content, salience, and price extraction suggestions.", "Shared", "Changing it affects the whole suggestion family."),
        ("Suggestion engine", "suggestion_travel", "Spots planned but unbooked travel and booking timing.", "Shared", "Prompt-only."),
        ("Suggestion engine", "suggestion_content", "Suggests content topics from RSS/newsletter signals.", "Shared", "Prompt-only."),
        ("Suggestion engine", "suggestion_opportunity", "Judges whether opportunity signals matter.", "Shared", "Prompt-only."),
        ("Suggestion engine", "salience_search_plan", "Plans semantic searches for signal relevance.", "Shared", "Prompt-only."),
        ("Suggestion engine", "travel_price_extract", "Extracts flight prices from alert emails.", "Shared", "Prompt-only."),
        ("Daily and weekly reports", "work_daily_brief", "Parent model for work daily brief calls.", "Shared", "Yesterday recap, today line, project signals."),
        ("Daily and weekly reports", "work_brief_recap", "Summarises yesterday's emails and meetings.", "Shared", "Prompt-only."),
        ("Daily and weekly reports", "work_brief_today", "Summarises today's calendar in one sentence.", "Shared", "Prompt-only."),
        ("Daily and weekly reports", "work_brief_project_salience", "Links current signals to project evidence.", "Shared", "Prompt-only."),
        ("Daily and weekly reports", "weekly_digest", "Writes the Sunday weekly digest sections.", "Shared", "Uses chats, emails, CRM."),
        ("Debrief", "debrief_interviewer", "Conducts the voice debrief.", "Per user", "Must be fast and concise."),
        ("Debrief", "debrief_extractor", "Extracts people, projects, and actions from debrief transcript.", "Per user", "Do not infer unstated actions."),
        ("Debrief", "debrief_transcriber", "Speech-to-text for spoken debrief answers.", "Shared", "Must be an STT model."),
        ("Debrief", "debrief_tts", "Text-to-speech for interviewer replies.", "Shared", "Must be a TTS model."),
        ("Debrief", "debrief_tts_voice", "Voice name passed to TTS.", "Shared", "Prompt-only voice preset."),
        ("LinkedIn pipeline", "linkedin_planner", "Generates search queries for a topic.", "Per user", "Recency-oriented."),
        ("LinkedIn pipeline", "linkedin_synthesiser", "Writes research briefing from sources.", "Per user", "Evidence-led."),
        ("LinkedIn pipeline", "linkedin_drafter", "Writes initial teaser post.", "Per user", "Short professional post."),
        ("LinkedIn pipeline", "linkedin_scorer", "Scores the draft against the rubric.", "Per user", "JSON scoring."),
        ("LinkedIn pipeline", "linkedin_carousel", "Generates carousel slide content.", "Per user", "Structured JSON."),
        ("LinkedIn pipeline", "linkedin_refiner", "Improves the teaser post.", "Per user", "Surgical editor."),
        ("LinkedIn pipeline", "linkedin_carousel_reviewer", "Reviews carousel JSON.", "Per user", "Quality gate."),
        ("LinkedIn pipeline", "linkedin_image", "Writes image-generation prompt.", "Per user", "No text in image."),
        ("LinkedIn pipeline", "linkedin_title", "Creates public display title when publishing.", "Per user", "Feeds portfolio/llms context."),
        ("LinkedIn tone", "spiciness_challenging_*", "Adds a more direct editorial stance.", "Per user", "Prompt-only modifier."),
        ("LinkedIn tone", "spiciness_provocative_*", "Adds bolder hot-take framing.", "Per user", "Prompt-only modifier."),
        ("Workday", "workday_narrative", "Converts a workday voice transcript into Markdown.", "Per user", "Used by Workday ingest."),
        ("Portfolio", "portfolio_chat", "Public Ask Me chat model.", "Per user", "Cheap, safe public model."),
        ("Portfolio", "jd_analyser", "Public job-description analyser.", "Per user", "Cheap, safe public model."),
        ("Knowledge layer", "embeddings", "Embeds documents, emails, facts, meetings for semantic retrieval.", "Shared", "Must be an embedding model; changing it re-indexes over time."),
        ("Knowledge layer", "atom_extractor", "Extracts durable claims from raw sources.", "Shared", "Nightly synthesis."),
        ("Knowledge layer", "completed_task_atom_extractor", "Promotes durable completed tasks into knowledge.", "Shared", "Nightly synthesis."),
        ("Knowledge layer", "entity_linker", "Links atoms to contacts, companies, or projects.", "Shared", "Handles ambiguity."),
        ("Knowledge layer", "cross_entity_synthesis", "Writes insight atoms across entities.", "Shared", "Patterns, gaps, opportunities."),
        ("Knowledge layer", "live_thread_synthesis", "Finds recurring ideas across sources without forcing CRM buckets.", "Shared", "Thread-level synthesis."),
        ("Knowledge layer", "interest_synthesis", "Names active work topics from meetings and calendar.", "Shared", "Feeds radar/brief."),
        ("CRM knowledge engine", "crm_source_triage", "Decides whether a source deserves CRM synthesis.", "Shared", "First stage of CRM loop."),
        ("CRM knowledge engine", "crm_duplicate_review", "Reviews duplicates and supersessions.", "Shared", "Prevents repeated facts/tasks."),
        ("CRM knowledge engine", "crm_action_projection", "Decides whether candidate actions become tasks/reminders.", "Shared", "High-confidence actions only."),
        ("CRM reports", "project_report", "Writes project reports from project evidence.", "Shared", "Use for report page."),
        ("CRM reports", "knowledge_query", "Answers Ask the Hub questions.", "Shared", "Atoms plus semantic search."),
        ("Wiki", "wiki_page_writer", "Converts documents and Q&A into wiki pages.", "Shared", "Structured knowledge pages."),
        ("Wiki", "wiki_image_vision", "Describes images before wiki page creation.", "Shared", "Vision model."),
        ("Newsletter intelligence", "newsletter_extractor", "Extracts topics from newsletter emails.", "Shared", "Runs on every newsletter."),
        ("Newsletter intelligence", "newsletter_briefing", "Writes weekly intelligence briefings.", "Per user", "Stronger writer recommended."),
        ("Writing tools", "ai_humanizer", "Removes AI-writing tells without inventing facts.", "Shared", "Keep output schema intact."),
        ("Nakai intelligence", "nakai_daily_briefing", "Writes daily regulatory PDF briefing.", "Shared", "Source-pack only."),
        ("Nakai intelligence", "nakai_ref_extraction", "Extracts substantive regulatory page content.", "Shared", "Reference-source pipeline."),
        ("Nakai intelligence", "nakai_ref_synthesis", "Compiles audit-horizon reference knowledge.", "Shared", "Feeds daily briefing."),
    ]


def add_model_slot_reference(doc):
    groups = {}
    for group, slot, purpose, scope, note in model_slots():
        groups.setdefault(group, []).append((slot, purpose, scope, note))
    for group, rows in groups.items():
        add_heading(doc, group, 2)
        add_table(
            doc,
            ["Slot", "What it does", "Scope", "User note"],
            rows,
            [2600, 4000, 1000, 1760],
            font_size=8.2,
        )


def main():
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("McLellan Hub User Manual")
    set_run_font(run, size=26, color="000000", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Hub, Admin, Prompt Library, Wiki, CRM, intelligence, and user-facing AI controls")
    set_run_font(r, size=12, color="555555")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Version 1.0 - Current as of 5 July 2026 - Private end-user reference")
    set_run_font(r, size=10, color="666666", italic=True)

    add_callout(
        doc,
        "Reader promise",
        "This is a user manual, not a developer or GitHub manual. It explains what each page, model, prompt, tool, function, purpose, and layer is for, when to use it, and what to expect from it.",
    )

    add_heading(doc, "How To Use This Manual", 1)
    add_para(doc, "Use the first half as a page-by-page operating guide. Use the admin chapters when you need to change behaviour, inspect what happened, or recover from a failed automation. Use the appendices when you need to understand a model slot, prompt, source type, or knowledge layer term.")
    add_bullets(doc, [
        "If you are doing everyday work, start with Hub Home, AI Chat, CRM, Tasks, Wiki, and Ask the Hub.",
        "If a background feature behaved oddly, check Admin: Knowledge, Jobs, Chat Logs, Newsletter Ingestion, Models, and Prompts.",
        "If you are changing model behaviour, change a visible model slot or prompt override rather than assuming one chat model controls everything.",
        "If a fact looks wrong, open its source before deleting it. Most Hub knowledge is source-backed and can be corrected at the source, by marking wrong, or through review.",
    ])

    add_heading(doc, "The Hub In One Page", 1)
    add_para(doc, "The Hub is a private working environment and second brain. It combines chat, projects, CRM, tasks, meetings, documents, email, wiki, flights, intelligence, content creation, and admin controls. The important idea is that most pages are views over source evidence and compiled knowledge, not isolated notebooks.")
    add_table(doc, ["Layer", "What it contains", "User-facing surfaces"], [
        ("Raw sources", "Emails, AgentMail, meetings, documents, tasks, chat messages, calendar events, RSS/newsletter stories, flights, debriefs, wiki pages.", "Source viewers, meeting intake, document uploads, newsletter ingestion, flights, debrief admin."),
        ("Extraction", "AI reads raw sources and proposes summaries, labels, actions, topics, attendees, facts, or claims.", "Meeting Intake, email taxonomy, newsletter ingestion, LinkedIn pipeline, prompt/model test arena."),
        ("Synthesis", "The Hub compares new evidence with what it already knows, handles duplicates/supersession, links entities, and decides whether actions matter.", "Knowledge review, Ask the Hub, CRM pages, project reports."),
        ("Compiled knowledge", "Source-backed atoms, insight atoms, thread themes, interest radar, project/contact/company knowledge, task projections.", "Person/company/project Knowledge panels, Ask the Hub, daily/weekly briefings."),
        ("Views and actions", "Human-facing pages and buttons for work: chat, tasks, CRM, wiki, content, intelligence, admin.", "Everything you use day to day."),
    ], [1800, 4100, 3460])

    add_heading(doc, "Daily Operating Rhythm", 1)
    add_numbered(doc, [
        "Open Hub Home to see current context, recent projects, and shortcuts.",
        "Use AI Chat for active work. Use a project chat when the work should accumulate memory or use project documents.",
        "Before meetings, open the person, company, or project page and review tasks, facts, knowledge, and recent meetings.",
        "After meetings, use Meeting Intake or the debrief flow so the Hub can capture actions and source-backed knowledge.",
        "Work from Tasks for operational commitments. Correct bad tasks with Wrong rather than silently ignoring them.",
        "Use Ask the Hub or Wiki Search when you need retrieval across the private knowledge base.",
        "Use Admin only when you need to inspect, correct, configure, or test the system.",
    ])

    add_heading(doc, "Page Inventory", 1)
    add_para(doc, "This inventory covers the user-visible Hub, CRM, Wiki, Newsletter, LinkedIn/content, Prompt Library, and admin entry points found in the current application.")
    add_table(doc, ["Route", "Page", "Purpose", "Use it when"], page_inventory(), [1650, 2100, 3650, 1960], font_size=8.2)

    doc.add_page_break()
    add_heading(doc, "Core User Pages", 1)

    add_heading(doc, "Hub Home", 2)
    add_para(doc, "Hub Home is the launch surface. It is not where knowledge is edited. It is where you orient yourself, jump back into recent work, and notice current signals such as today's calendar, recent conversations, active projects, newsletter activity, and token burn.")
    add_bullets(doc, [
        "Use recent conversations to resume work without searching logs.",
        "Use project tiles when work belongs to an ongoing stream.",
        "Use the navigation tiles to reach CRM, Wiki, Intelligence, Content, Tasks, Flights, Debrief, and Admin.",
    ])

    add_heading(doc, "AI Chat And Project Chat", 2)
    add_para(doc, "AI Chat is the general assistant. It can answer, draft, analyse files, export results, recall prior conversations, and save useful outputs. Project Chat is the same style of assistant, but scoped to a project and its saved memory/documents.")
    add_table(doc, ["Control or command", "What it does", "Good use"], [
        ("Model picker", "Chooses the model for this chat response.", "Use stronger models for synthesis, long documents, or high-stakes drafting."),
        ("Search mode", "Allows web/search-assisted work where configured.", "Use for current facts, product/law/news questions, and evidence-led research."),
        ("Upload", "Extracts and analyses files. Project uploads can become reusable project documents.", "Use for PDFs, Word, PowerPoint, Excel, text, markdown, CSV, HTML, EPUB, and images."),
        ("Save to project", "Keeps an answer as project context.", "Use when a response should be reusable in the project."),
        ("Save to Wiki", "Creates a durable wiki page from a useful Q&A or document.", "Use when the answer should be findable outside the chat."),
        ("Export", "Produces a Word/PDF/Google Docs style output where supported.", "Use for deliverables."),
        ("Ratings", "Marks an answer as good or poor.", "Use to help later model evaluation."),
        ("/recall terms", "Searches old chats and can inject recalled context.", "Use when you remember the topic but not the conversation."),
        ("/tasks add title", "Creates a Google Task.", "Use for quick capture from chat."),
        ("/tasks done words", "Completes the first matching task.", "Use for quick cleanup."),
        ("/slug message", "Works inside or with the named project.", "Use to target a project from normal chat."),
    ], [2100, 3900, 3360])

    add_heading(doc, "CRM And Relationship Pages", 2)
    add_para(doc, "The CRM is a briefing system for relationships and workstreams. It is not just an address book. People, companies, projects, meetings, tasks, facts, and knowledge claims are connected through source evidence and synthesis.")
    add_table(doc, ["CRM surface", "What to look for", "Normal action"], [
        ("People", "Contacts and search.", "Open a person before or after an interaction."),
        ("Person page", "Timeline, tasks, companies, projects, meetings, knowledge claims, edit/merge controls.", "Review before calls; add task or fact; mark bad facts wrong."),
        ("Companies", "Organisations and types.", "Find or create companies."),
        ("Company page", "People, tasks, knowledge, meetings, timeline.", "Prepare for organisation-level work."),
        ("Meetings", "Meeting list and creation.", "Find meeting history or create a manual record."),
        ("Meeting Intake", "Transcript upload/paste and extracted outputs.", "Process meeting evidence into summaries, actions, projects, and CRM."),
        ("Tasks", "Google-backed task queue with Hub context.", "Work the queue; sync; show history; restore local deletes."),
        ("Reminders", "Reminder states and acknowledgement actions.", "Handle nudges that are not full tasks."),
        ("Projects", "Project list and project pages.", "Understand a workstream from tasks, people, docs, meetings, and knowledge."),
        ("Project Report", "AI-generated report from evidence.", "Create a progress narrative or update."),
        ("Ask the Hub", "Free-form questions over compiled knowledge.", "Ask relationship or project questions across sources."),
    ], [2200, 3900, 3260])

    add_heading(doc, "Tasks", 2)
    add_para(doc, "Tasks are backed by Google Tasks, with Hub-only context layered on top. Google stores the task and due date. The Hub stores extra links such as person, company, project, source, and due time.")
    add_bullets(doc, [
        "Done completes the task in Google Tasks and the Hub cache.",
        "Delete hides a task locally; it does not necessarily delete it from Google.",
        "Restore removes the local deletion marker.",
        "Wrong teaches the Hub that a created task was a bad projection and should not be repeated.",
        "Show history reveals completed or locally deleted context.",
        "Due time is Hub-only because Google Tasks itself stores dates, not times.",
    ])

    add_heading(doc, "Meeting Intake", 2)
    add_para(doc, "Meeting Intake turns a transcript into structured evidence. It can produce a meeting summary, attendees, CRM updates, action register, project notes, open questions, and warnings. It is strongest when you provide a title hint, date, project hint, and clean transcript.")
    add_table(doc, ["Field", "What to enter", "Why it matters"], [
        ("Title hint", "A human-readable meeting name.", "Helps title and context matching."),
        ("Date", "Meeting date.", "Keeps timeline correct."),
        ("Attach to meeting", "Existing meeting if this transcript belongs to it.", "Avoids duplicate meeting records."),
        ("Project", "Known project if applicable.", "Gives the extractor a safe anchor."),
        ("Transcript upload/paste", "Full transcript text.", "This is the source evidence."),
        ("Identify speakers", "Map generic speaker labels to real people.", "Prevents false person records and bad ownership."),
        ("Save to CRM", "Commits the reviewed extraction.", "Turns intake into usable CRM/project/task context."),
    ], [1900, 3500, 3960])

    add_heading(doc, "Wiki", 2)
    add_para(doc, "The Wiki is the durable, human-readable knowledge surface. It holds authored pages, imported material, source-backed pages, person/project notes, and searchable vault content. Wiki Search can also synthesize an answer from matched private sources.")
    add_bullets(doc, [
        "Use Search first when you have a question.",
        "Use Browse when you want to inspect indexed material by type.",
        "Use Add to queue web or YouTube URLs for ingestion.",
        "Use Graph and Orphans as maintenance views when the knowledge base feels disconnected.",
        "Use source links and related pages to follow evidence instead of trusting a detached summary.",
    ])

    add_heading(doc, "Newsletter Intelligence", 2)
    add_para(doc, "Newsletter Intelligence turns incoming newsletter emails and RSS/creator feeds into topics, interests, briefings, and publishable source-backed knowledge. It separates raw ingestion from briefing generation so you can audit what was extracted before using it.")
    add_table(doc, ["Area", "Purpose", "Main actions"], [
        ("Topics and interests", "Review extracted topics and decide what matters.", "Toggle topics, auto interests, sync labels, suggest interests."),
        ("Briefings", "Create, preview, send, publish, delete, export PDF, save to wiki, or send to NoteMax.", "Generate and curate weekly or creator-specific briefings."),
        ("Sources", "Maintain source priority and schedules.", "Set priority, enable/disable, add schedules."),
        ("Creators", "Follow creator/RSS feeds, read articles, group stories, generate briefings.", "Fetch, read, brief, toggle, delete."),
    ], [2200, 3800, 3360])

    add_heading(doc, "LinkedIn Content", 2)
    add_para(doc, "The Content tool is a pipeline, not a text box. It researches a topic, drafts a short teaser, scores it, refines it, creates carousel content, writes an image prompt, and tracks queue/published state.")
    add_numbered(doc, [
        "Start on Content Create with a topic or source URL.",
        "Choose tone: Professional, Challenging, or Provocative.",
        "Let the planner and synthesiser gather and analyse current sources.",
        "Review the draft, score, refiner output, carousel, and image prompt.",
        "Use Queue for processing, draft, scheduled, or error items.",
        "Use Published to audit live content and sort/filter history.",
        "Use Plan to set cadence, research topics by day, and manage suggestions.",
    ])

    add_heading(doc, "Daily Debrief", 2)
    add_para(doc, "The debrief is a voice-led end-of-day operational capture. It asks about unplanned work, then works through known calendar/email/work topics. It is designed for driving or hands-free use, so replies are short and one question at a time.")
    add_bullets(doc, [
        "It should capture what happened, decisions, blockers, and next actions.",
        "It should not become therapy or general reflection unless you explicitly steer it there.",
        "The extractor only uses what you explicitly said, especially for actions.",
        "Admin Debrief lets you review transcripts and extracted outputs.",
    ])

    add_heading(doc, "Flights", 2)
    add_para(doc, "Flights records travel history, imports bookings, looks up flight status, and supports travel-related task preparation. Ryanair booking emails can also create flight records through email processing.")
    add_bullets(doc, [
        "Use manual add for a flight that was not imported.",
        "Use file import for booking documents.",
        "Use lookup or bulk lookup to refresh details where the API supports it.",
        "Use the statistics area for flight history and travel totals.",
    ])

    doc.add_page_break()
    add_heading(doc, "Admin Guide", 1)
    add_para(doc, "Admin pages are for configuration, audit, and recovery. They are user-facing, but they can change background behaviour immediately. When unsure, inspect before editing.")
    add_table(doc, ["Admin route", "Page", "Purpose"], admin_inventory(), [2300, 2600, 4460], font_size=8.6)

    add_heading(doc, "Models Admin", 2)
    add_para(doc, "Models Admin has three separate jobs: chat models, system model slots, and model tiers. The chat model selected in AI Chat affects that conversation. System model slots affect background features such as email classification, meeting intake, Ask the Hub, debrief, newsletter extraction, and knowledge synthesis.")
    add_table(doc, ["Admin tab", "What it controls", "What to watch"], [
        ("Models", "Available model records: key, label, provider endpoint, model id, tier, search mode, enabled state, costs, context length.", "Do not delete a model that a system slot depends on unless you intend it to fall back."),
        ("System", "Which model each background slot uses.", "Changes apply immediately. Per-user slots affect only the current user; system slots affect everyone."),
        ("Prompts", "The prompt text for visible slots.", "Edited prompts override built-in defaults. Restore default clears the override."),
        ("Tiers", "Chat picker grouping and default search mode by tier.", "Changing tier keys after use can confuse model records."),
    ], [2100, 4500, 2760])

    add_heading(doc, "Prompt Overrides", 2)
    add_para(doc, "Every visible prompt slot has a built-in default in the application and may have an admin override. If a prompt is edited, the edited text is what runs. Restore default removes the override.")
    add_bullets(doc, [
        "Dynamic placeholders such as [DATE], [NAME], [CALENDAR], [CONTEXT], [NOTE], [FIELD], and [SOURCE_TEXT] are filled at runtime. Do not remove them unless you want to change what context the model receives.",
        "Shape for model proposes a rewrite in the style of the assigned model family. It does not save automatically. Review before saving.",
        "Prompt-only slots change instructions appended to or used by a parent model; they do not select a model themselves.",
        "When testing a changed prompt, use Admin Test or the smallest affected workflow first.",
    ])

    add_heading(doc, "Knowledge Review", 2)
    add_para(doc, "Knowledge Review is where the Hub shows what the synthesis layer is unsure about or has recently decided. It is the most important admin page for keeping the second brain trustworthy.")
    add_table(doc, ["Item", "Meaning", "Normal action"], [
        ("Review queue", "Contradictions, stale facts, duplicate contacts, or atoms needing approval/rejection.", "Approve, reject, mark status, or investigate source."),
        ("Knowledge receipts", "Audit trail of prompt decisions and payloads.", "Use to see why a model acted or did not act."),
        ("Dedup action", "Retroactive merge of near-duplicate atoms.", "Run when duplicates are visible, then inspect results."),
        ("Atom status", "Active, stale, wrong, rejected, or otherwise statused claims.", "Prefer marking wrong/rejected over deleting evidence."),
    ], [2200, 4200, 2960])

    add_heading(doc, "Jobs And Connectivity", 2)
    add_para(doc, "Jobs are scheduled or queued background work. Connectivity is the health view for isolated records. If something did not happen, check Jobs first, then the relevant admin audit page.")
    add_bullets(doc, [
        "Pending jobs are work waiting to run. Recent jobs show completed, failed, or cancelled items.",
        "Cancel only pending jobs you are sure are no longer wanted.",
        "Connectivity highlights orphaned contacts, projects, documents, tasks, flights, or facts that need synthesis or cleanup.",
        "Run mycelium/connectivity when new material exists but expected links have not appeared yet.",
    ])

    add_heading(doc, "Email Taxonomy And Ingestion", 2)
    add_para(doc, "The email taxonomy is how the Hub keeps Gmail and AgentMail organised for human navigation and downstream synthesis. In normal operation, email ingestion stores source evidence and lets the CRM knowledge engine decide durable facts/actions.")
    add_bullets(doc, [
        "Canonical labels should be clear, stable, and useful in Gmail.",
        "AgentMail review items remain visible until classified or learned.",
        "Trigger Email Fetch or AgentMail Fetch when you need to process immediately instead of waiting for the schedule.",
        "Use Newsletter Ingestion to inspect newsletter extraction success, fallback, and source data.",
    ])

    add_heading(doc, "Admin Test Arena", 2)
    add_para(doc, "The test arena is where you compare models and prompts before trusting a change. It can run a prompt on one or more models, attach a file, improve a prompt, and save logs.")
    add_numbered(doc, [
        "Write the prompt or choose a saved test.",
        "Select one or more models or tiers.",
        "Attach a file only if the task requires it.",
        "Use Improve prompt when the instruction is rough.",
        "Run and compare outputs for accuracy, format compliance, source discipline, and cost.",
        "Save useful test cases so future model changes can be compared.",
    ])

    doc.add_page_break()
    add_heading(doc, "Model And Prompt Slot Reference", 1)
    add_para(doc, "This reference explains the named AI controls in user terms. Shared slots are system-wide. Per-user slots are stored for the current Hub user. Prompt-only slots alter instructions but inherit a parent model.")
    add_model_slot_reference(doc)

    doc.add_page_break()
    add_heading(doc, "Tools, Functions, And Layers", 1)
    add_heading(doc, "User-Facing Tools", 2)
    add_table(doc, ["Tool", "Purpose", "Main inputs", "Main outputs"], [
        ("AI Chat", "General assistant workspace.", "Question, model, search mode, files, project context.", "Answer, export, saved project memory, wiki page, tasks."),
        ("Project Documents", "Reusable source material for project chat and synthesis.", "Uploaded files or Drive import.", "Markdown extraction, tasks, wiki pages, project context."),
        ("Meeting Intake", "Turn transcripts into source-backed relationship/project evidence.", "Transcript, title/date/project hints, speaker map.", "Meeting summary, actions, CRM updates, project notes."),
        ("Google Tasks", "Operational task system.", "Manual tasks, projected actions, chat commands.", "Google Tasks plus Hub links/source metadata."),
        ("Ask the Hub", "Question-answering over compiled knowledge.", "Free-form question.", "Answer with evidence context and relevant insights."),
        ("Wiki", "Durable knowledge base.", "Manual pages, saved chat answers, documents, ingested URLs.", "Searchable pages, graph, related links."),
        ("Newsletter Intelligence", "Extract topics and write briefings.", "Newsletter/RSS sources, interests, schedules.", "Topics, briefings, PDFs, wiki pages, published items."),
        ("LinkedIn Content", "Research and prepare professional content.", "Topic, source URL, tone.", "Draft, score, refined post, carousel, image prompt."),
        ("Debrief", "Hands-free end-of-day operational capture.", "Voice answers and known daily context.", "Transcript, actions, CRM/project evidence."),
        ("Flights", "Travel records and flight intelligence.", "Booking emails, manual entries, imports, lookup.", "Flight history, status, stats, prep tasks."),
        ("Prompt Library", "Reusable prompt assets and optimization.", "Prompt text, examples, purpose, target model.", "Saved prompts, adapted prompts, optimized winners, agent packs."),
    ], [2000, 3100, 3000, 1260], font_size=8.4)

    add_heading(doc, "Knowledge Terms", 2)
    add_table(doc, ["Term", "User meaning"], [
        ("Source evidence", "The original item the Hub read: email, meeting, document, task, chat, RSS story, or wiki page."),
        ("Embedding", "A semantic index entry that lets the Hub find related text even when the words do not match exactly."),
        ("Atom", "A source-backed claim such as a relationship, address, role, decision, need, or project fact."),
        ("Insight atom", "A synthesized pattern, gap, opportunity, or connection across multiple atoms or entities."),
        ("Entity linker", "The step that decides which person, company, or project a claim is about."),
        ("Duplicate review", "The step that decides whether a new source confirms, updates, supersedes, or duplicates existing knowledge."),
        ("Action projection", "The step that decides whether a possible action should become a Google Task or reminder."),
        ("Knowledge receipt", "An audit record of what a prompt/model saw and decided."),
        ("Stale claim", "A claim that is old or unconfirmed. It is hidden from default view but remains searchable and can be revived by new evidence."),
        ("Wrong claim/task", "A user-marked error. Marking wrong is useful because it teaches the system what not to repeat."),
        ("Mycelium", "The connection pass that looks for orphaned data and grows useful links/tasks across modules."),
        ("Interest radar", "A daily synthesis of active work topics from meetings and calendar, used by briefings and suggestions."),
        ("Live thread", "A recurring idea noticed across sources without forcing it into a single person/project bucket."),
    ], [2300, 7060], font_size=8.6)

    add_heading(doc, "Source Kinds", 2)
    add_table(doc, ["Source kind", "What it means", "Where you see it"], [
        ("email_summary", "Gmail or AgentMail summary stored as source evidence.", "CRM sources, knowledge claims, receipts."),
        ("meeting_intake", "Transcript-derived meeting evidence.", "Meeting pages, project pages, knowledge claims."),
        ("document", "Uploaded or Drive-derived document extraction.", "Project documents, wiki pages, task extraction."),
        ("open_task", "Current operational task state.", "Tasks, duplicate review context."),
        ("completed_task", "Completed task considered for durable knowledge.", "Knowledge atoms and project/contact history."),
        ("crm_fact", "Manual or legacy curated CRM fact.", "Timelines and synthesis context."),
        ("newsletter/RSS story", "Incoming intelligence source.", "Newsletter, content suggestions, live threads."),
        ("wiki page", "Durable authored or generated knowledge page.", "Wiki search, project context, chat context."),
    ], [2200, 4300, 2860], font_size=8.5)

    add_heading(doc, "What To Do When Something Looks Wrong", 1)
    add_table(doc, ["Symptom", "First place to check", "Likely fix"], [
        ("A fact on a person/company/project is wrong.", "Open the source link from the Knowledge panel.", "Mark wrong/reject in page or Admin Knowledge; correct the source if needed."),
        ("A task should not have been created.", "Open task detail and inspect source.", "Use Wrong, then review task learning/receipts."),
        ("A meeting created bad actions.", "Meeting Intake record and speaker map.", "Correct speaker mapping; delete or mark bad outputs; rerun/save carefully."),
        ("Newsletter briefing missed items.", "Admin Newsletter Ingestion and source priorities.", "Check extraction status, source priority, interest toggles, schedules."),
        ("A model output changed style or quality.", "Admin Models/System and Prompts.", "Check slot assignment, prompt override, and test arena logs."),
        ("Ask the Hub gives stale or incomplete answer.", "Admin Knowledge, embeddings status, source evidence.", "Check whether evidence exists, atoms are active, or search is indexed."),
        ("A background action did not run.", "Admin Jobs and Connectivity.", "Inspect failed/pending jobs; run connectivity/mycelium if links are missing."),
        ("A page seems to lack context.", "Project admin, documents, wiki tags, source links.", "Import/link the right documents or let synthesis run."),
    ], [2700, 3300, 3360], font_size=8.4)

    add_heading(doc, "Safe Admin Change Checklist", 1)
    add_numbered(doc, [
        "Identify the exact surface affected: chat, meeting intake, newsletter, debrief, knowledge, content, portfolio, or reports.",
        "Find the visible model slot or prompt slot for that surface.",
        "Read the existing prompt and note required output format or placeholders.",
        "Test with a known example in Admin Test or the relevant workflow.",
        "Save the smallest useful change.",
        "Inspect logs, receipts, or output for the next real run.",
        "If the change creates worse output, restore default or use fallback.",
    ])

    add_heading(doc, "Glossary", 1)
    add_table(doc, ["Word", "Plain-English meaning"], [
        ("Hub", "The private personal workspace at dchat/nchat."),
        ("Admin", "Configuration, audit, and recovery pages for the Hub."),
        ("Project", "A continuing workstream with chat memory, documents, CRM links, and knowledge."),
        ("Prompt", "The instruction text a model follows for a named task."),
        ("Model slot", "A named setting that chooses which model a background feature uses."),
        ("Fallback", "The built-in model or prompt used when no admin override exists."),
        ("Per-user", "Stored separately for Douglas/Nakai or the current user."),
        ("System", "Shared across users."),
        ("Receipt", "Audit record of an AI decision."),
        ("Compiled knowledge", "The Hub's source-backed interpretation of raw material."),
        ("Raw source", "The original thing ingested before interpretation."),
        ("Synthesis", "The process of re-reading sources and connecting them into useful knowledge."),
    ], [2500, 6860], font_size=8.6)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
